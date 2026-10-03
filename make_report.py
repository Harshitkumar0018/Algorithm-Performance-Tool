"""Builds Final_Report.docx from the CSVs, graphs and screenshots."""
import pandas as pd
from docx import Document
from docx.shared import Inches

search = pd.read_csv("reports/search_results.csv")
fact = pd.read_csv("reports/factorial_results.csv")
loops = pd.read_csv("reports/loops_results.csv")
cx = pd.read_csv("reports/complexity_table.csv")

doc = Document()


def table(df):
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)


def para(text):
    doc.add_paragraph(text)


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def at(df, name, size, col="Time_ms"):
    return float(df[(df.Algorithm == name) & (df.Size == size)][col].iloc[0])


doc.add_heading("Algorithm Performance Measurement and Benchmarking Tool", 0)
para("Lab Assignment 2 - Final Report")

doc.add_heading("1. Introduction", 1)
doc.add_heading("Importance of algorithm analysis", 2)
para("Algorithm analysis tells us how the running time and memory of a program grow as the input "
     "grows. Two programs that give the same answer can differ enormously in cost, so choosing "
     "the right one matters before deployment.")
doc.add_heading("Need for performance measurement", 2)
para("Theory (Big-O) describes growth but hides constants, hardware, and language overhead. "
     "Measuring real execution time and memory lets us check theory against practice and see "
     "where an algorithm that is fine on small data becomes too slow on large data.")
para("This project is a Streamlit tool (app.py) with a benchmarking engine (benchmark_engine.py), "
     "a search module (search_analysis.py) and code snippets (code_benchmark.py). Time is measured "
     "with time.perf_counter (average of several runs) and peak memory with tracemalloc.")

doc.add_heading("2. Experimental Results", 1)
doc.add_heading("Search benchmark (worst case: last element)", 2)
table(search.round(4))
doc.add_picture("graphs/search.png", width=Inches(5))
doc.add_heading("Factorial benchmark", 2)
table(fact.round(4))
doc.add_picture("graphs/factorial.png", width=Inches(5))
doc.add_heading("Loop benchmark", 2)
para("The nested loop performs n x n work, so its largest size is 5,000 (25 million operations) "
     "instead of 50,000 (2.5 billion), which would take many minutes.")
table(loops.round(4))
doc.add_picture("graphs/loops.png", width=Inches(5))
doc.add_heading("Complexity table: theoretical vs observed", 2)
table(cx)
doc.add_heading("Screenshots of outputs", 2)
for p in ["search_output", "code_output", "complexity_table"]:
    doc.add_picture(f"screenshots/{p}.png", width=Inches(6))

doc.add_heading("3. Analysis", 1)
doc.add_heading("Comparison of search algorithms", 2)
ls, bs = at(search, "Linear Search", 50000), at(search, "Binary Search", 50000)
para(f"At 50,000 elements linear search took {ls:.3f} ms and made 50,000 comparisons, while binary "
     f"search took {bs:.4f} ms and made 16 comparisons (about {ls / bs:.0f}x faster). Linear "
     "search time grows in proportion to n; binary search time is nearly flat, matching O(n) "
     "and O(log n). Binary search needs sorted data, which is the price of its speed.")
doc.add_heading("Recursive vs iterative implementations", 2)
rt, it = at(fact, "Recursive Factorial", 1000), at(fact, "Iterative Factorial", 1000)
rm, im = at(fact, "Recursive Factorial", 1000, "Memory_KB"), at(fact, "Iterative Factorial", 1000, "Memory_KB")
para(f"For n=1000 the recursive version took {rt:.3f} ms and {rm:.1f} KB; the iterative version took "
     f"{it:.3f} ms and {im:.1f} KB. Both do n multiplications, so timing is similar (differences "
     "are small and noisy), but recursion uses roughly 10x more memory because each call adds a "
     "stack frame (O(n) space vs O(1) extra space for iteration).")
doc.add_heading("Single loop vs nested loop", 2)
sl, nl = at(loops, "Single Loop", 5000), at(loops, "Nested Loop", 5000)
para(f"At 5,000 iterations the single loop took {sl:.3f} ms and the nested loop {nl:.1f} ms "
     f"(about {nl / sl:.0f}x slower). Each 2x increase in n roughly quadruples the nested loop "
     "time, confirming O(n^2), while the single loop doubles, confirming O(n).")
doc.add_heading("Complexity analysis and observations", 2)
bullets([
    "The log-log slope of time vs size approximates the exponent: about 1 for linear search and the "
    "single loop, about 0 for binary search, and about 2 for the nested loop.",
    "Factorial slopes are above 1 (about 1.4-1.7) even though the loop count is O(n): the numbers "
    "become huge integers, so each multiplication itself costs more as n grows.",
    "Very small timings (binary search, ~0.01 ms) are noisy, so averages of several runs are used.",
    "Memory of search and loop programs is tiny because they store no new data structures.",
])
doc.add_heading("Which is better, and time-space trade-offs", 2)
bullets([
    "Search: binary search is better for large sorted data; linear search is fine for tiny or unsorted data.",
    "Factorial: iterative is better (same time, much less memory, no recursion limit).",
    "Loops: a single pass is far better; avoid nested loops when a single pass will do.",
    "Trade-off: recursion trades memory for simpler code; binary search trades a sorting step up front for fast lookups.",
])

doc.add_heading("4. Conclusion", 1)
doc.add_heading("Key findings", 2)
para("Measured results agree with theory: growth in input size affects O(n^2) code far more than "
     "O(n) code, and O(log n) code barely changes. Theory ignores costs such as big-integer "
     "arithmetic and stack memory, which measurement reveals.")
doc.add_heading("Most efficient implementations observed", 2)
bullets(["Binary Search (search)", "Iterative Factorial (factorial)", "Single Loop (loops)"])
doc.add_heading("Future enhancements", 2)
bullets([
    "Add sorting algorithms and more search techniques (jump, interpolation).",
    "Let users paste and benchmark their own code snippets.",
    "Add average/best-case inputs and confidence intervals over many runs.",
    "Export benchmark reports from the app as CSV or PDF.",
])
doc.add_heading("References", 2)
para("Horowitz, Sahni, Rajasekaran - Fundamentals of Computer Algorithms; Python, Streamlit, "
     "matplotlib and memory_profiler documentation.")

doc.save("Final_Report.docx")
